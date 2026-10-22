"""사다리 발표 자료 초안 생성기 (python-pptx).

디자인: 16:9 · Paperlogy 전용 · 헤더 위치 고정 · 따뜻한 종이색 캔버스 · 제목/구분 슬라이드는 딥 인디고.
실행: python docs/presentation/build_deck.py   → docs/presentation/_draft.pptx (폰트 미포함 초안)
이후 finalize_deck.ps1 이 PowerPoint로 Paperlogy 폰트를 파일에 포함해 최종본을 만든다.

이 스크립트는 '초안'을 만드는 용도다. 사용자 테스트 결과 등은 최종 .pptx를 PowerPoint에서 직접 고친다.
"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
SCREENS = HERE / "screens"
OUT = HERE / "_draft.pptx"

# ------------------------------------------------------------------ 디자인 토큰
C = {
    "canvas": "F6F5F4", "surface": "FFFFFF", "hairline": "E6E6E6",
    "ink": "111111", "ink2": "31302E", "muted": "615D59", "faint": "A39E98",
    "primary": "0075DE", "indigo": "213183",
    "sky": "62AEF0", "purple": "D6B6F6", "purple_deep": "391C57", "pink": "FF64C8",
    "orange": "DD5B00", "teal": "2A9D99", "green": "1AAE39", "white": "FFFFFF",
    "night_text": "D9DDF2", "night_muted": "A9B2DE",
    "lime": "9FE870", "lime_deep": "163300",  # 로고 심볼 (웹과 동일)
}
F = {
    "xbold": "Paperlogy 8 ExtraBold", "bold": "Paperlogy 7 Bold", "semi": "Paperlogy 6 SemiBold",
    "medium": "Paperlogy 5 Medium", "regular": "Paperlogy 4 Regular",
}
# 스타일: (굵기, 크기pt, 기본색, 자간(1/100pt))
S = {
    "display1": ("bold", 60, "ink", -150), "display2": ("bold", 50, "ink", -120),
    "h1": ("bold", 38, "ink", -60), "sub": ("regular", 16, "ink2", 0),
    "h2": ("bold", 24, "ink", -40), "h3": ("bold", 19, "ink", -20), "title": ("semi", 16, "ink", 0),
    "body": ("regular", 15, "ink2", 0), "body_sm": ("regular", 13, "ink2", 0),
    "caption": ("regular", 11.5, "muted", 0), "eyebrow": ("semi", 12, "primary", 10),
    "stat": ("xbold", 40, "ink", -60), "stat_xl": ("xbold", 96, "ink", -250),
    "button": ("medium", 15, "white", 0), "chip": ("semi", 12, "ink", 0), "code": ("regular", 9, "faint", 0),
}

X0, CW = 0.6, 12.13  # 안전 영역 왼쪽 x, 콘텐츠 폭 (인치)


# ------------------------------------------------------------------ 저수준 도우미
def _rgb(color: str) -> RGBColor:
    return RGBColor.from_string(C.get(color, color))


def style_run(run, style: str, color: str | None = None, size: float | None = None) -> None:
    weight, default_size, default_color, spc = S[style]
    family = F[weight]
    font = run.font
    font.name = family
    font.size = Pt(size or default_size)
    font.bold = False
    font.color.rgb = _rgb(color or default_color)
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "ko-KR")
    if spc:
        rpr.set("spc", str(spc))
    latin = rpr.find(qn("a:latin"))
    for tag in ("a:cs", "a:ea"):  # 한글은 ea 글꼴을 쓰므로 반드시 지정한다
        old = rpr.find(qn(tag))
        if old is not None:
            rpr.remove(old)
        latin.addnext(rpr.makeelement(qn(tag), {"typeface": family}))


def text(slide, x, y, w, h, paras, *, align="left", anchor="top", space_after=0, line=None):
    """paras: 문단 목록. 문단 = (글, 스타일[, 색]) 또는 그런 튜플의 리스트(한 문단 여러 런)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, tuple):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        p.space_after = Pt(space_after)
        if line:
            p.line_spacing = line
        for run_spec in (para if isinstance(para, list) else [para]):
            r = p.add_run()
            r.text = run_spec[0]
            style_run(r, run_spec[1], run_spec[2] if len(run_spec) > 2 else None)
    return tb


def _alpha(fill_parent, alpha: int) -> None:
    clr = fill_parent.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    clr.append(clr.makeelement(qn("a:alpha"), {"val": str(alpha * 1000)}))


def _soft_shadow(shape, strength: int = 1) -> None:
    sppr = shape._element.spPr
    eff = sppr.find(qn("a:effectLst"))
    blur, dist, alpha = (152400, 25400, 7000) if strength == 1 else (330200, 76200, 9000)
    sh = eff.makeelement(qn("a:outerShdw"), {"blurRad": str(blur), "dist": str(dist), "dir": "5400000",
                                               "algn": "t", "rotWithShape": "0"})
    clr = sh.makeelement(qn("a:srgbClr"), {"val": "000000"})
    clr.append(clr.makeelement(qn("a:alpha"), {"val": str(alpha)}))
    sh.append(clr)
    eff.append(sh)


def box(slide, x, y, w, h, *, fill="surface", line=None, radius=0.0, shadow=0, alpha=None,
        shape=None, dash=False, line_w=0.75):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius and kind in (MSO_SHAPE.ROUNDED_RECTANGLE, MSO_SHAPE.ROUND_2_SAME_RECTANGLE):
        s.adjustments[0] = min(radius / min(w, h), 0.5)
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = _rgb(fill)
        if alpha is not None:
            _alpha(s._element.spPr, alpha)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = _rgb(line)
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if shadow:
        _soft_shadow(s, shadow)
    return s


def label_in(shape, paras, *, align="center", anchor="middle", margin=0.08):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE}[anchor]
    if isinstance(paras, tuple):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER}[align]
        for spec in (para if isinstance(para, list) else [para]):
            r = p.add_run()
            r.text = spec[0]
            style_run(r, spec[1], spec[2] if len(spec) > 2 else None)


def _text_width(s: str, per_wide: float, per_narrow: float) -> float:
    return sum(per_wide if ord(ch) > 0x2E80 else per_narrow for ch in s)


def pill(slide, x, y, label, *, style="chip", fill="surface", color=None, line="hairline", h=0.34, pad=0.22,
         w=None, shadow=0):
    size = S[style][1]
    width = w or pad * 2 + _text_width(label, size / 72 * 1.0, size / 72 * 0.56)
    s = box(slide, x, y, width, h, fill=fill, line=line, radius=h / 2, shadow=shadow)
    label_in(s, (label, style, color), margin=0.04)
    return s, width


def dot(slide, x, y, d, color, alpha=None):
    return box(slide, x, y, d, d, fill=color, shape=MSO_SHAPE.OVAL, alpha=alpha)


def num_badge(slide, x, y, d, label, color, text_color="white"):
    s = dot(slide, x, y, d, color)
    label_in(s, (label, "chip", text_color), margin=0)
    return s


def arrow(slide, x1, y1, x2, y2, color="faint", width=1.5, both=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = _rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if both:
        ln.append(ln.makeelement(qn("a:headEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def phone(slide, img: Path, x, y, h):
    w = h * 750 / 1624
    box(slide, x - 0.09, y - 0.09, w + 0.18, h + 0.18, fill="1B1B1B", radius=0.3, shadow=2)
    pic = slide.shapes.add_picture(str(img), Inches(x), Inches(y), Inches(w), Inches(h))
    pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
    geom = pic._element.spPr.find(qn("a:prstGeom"))
    av = geom.find(qn("a:avLst"))
    if av is None:
        av = geom.makeelement(qn("a:avLst"), {})
        geom.append(av)
    av.append(av.makeelement(qn("a:gd"), {"name": "adj", "fmla": "val 9000"}))
    return w


# ------------------------------------------------------------------ 슬라이드 골격
def new_slide(prs, bg="canvas"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(bg)
    return slide


def header(slide, chapter: str, title: str, subtitle: str | None = None):
    """고정 헤더: 챕터 라벨(0.42") → 제목(0.86") → 부제(1.52"). 모든 콘텐츠 슬라이드에서 같은 위치."""
    pill(slide, X0, 0.42, chapter, style="eyebrow", h=0.32, pad=0.16)
    text(slide, X0, 0.86, CW, 0.62, (title, "h1"))
    if subtitle:
        text(slide, X0, 1.52, CW, 0.36, (subtitle, "sub"))


def footer(slide, section: str, page: int, dark=False):
    color = "night_muted" if dark else "faint"
    text(slide, X0, 7.0, 8, 0.25, (f"사다리  ·  {section}", "caption", color))
    text(slide, X0 + CW - 1.0, 7.0, 1.0, 0.25, (str(page), "caption", color), align="right")


def logo_placeholder(slide, x, y, dark=True):
    """로고 심볼(A안: 위로 길어지는 가로대 세 칸)은 확정. 워드마크 서체는 미정이라 자리만 표시한다."""
    d = 0.46
    u = d / 40  # 웹 SVG와 같은 40×40 격자
    box(slide, x, y, d, d, fill="lime", radius=11 * u)
    for bx, by, bw in ((10, 26, 10), (10, 17.5, 15), (10, 9, 20)):
        box(slide, x + bx * u, y + by * u, bw * u, 5 * u, fill="lime_deep", radius=2.5 * u)
    s = box(slide, x + d + 0.14, y, 1.9, d, fill=None, line="night_muted" if dark else "faint", radius=0.1, dash=True)
    label_in(s, ("워드마크 · 서체 미정", "caption", "night_muted" if dark else "faint"))


def constellation(slide, x, y, w, h):
    """딥 인디고 슬라이드 장식: 스티커 색 가로대를 가진 사다리 + 별."""
    rail = 0.13
    left, right = x + w * 0.22, x + w * 0.78 - rail
    box(slide, left, y, rail, h, fill="white", alpha=35, radius=rail / 2)
    box(slide, right, y, rail, h, fill="white", alpha=35, radius=rail / 2)
    rungs = ["purple", "pink", "orange", "teal", "green", "sky"]
    gap = h / (len(rungs) + 0.6)
    for i, color in enumerate(rungs):
        ry = y + h - gap * (i + 0.9)
        box(slide, left - 0.05, ry, right - left + rail + 0.1, 0.17, fill=color, radius=0.085)
    stars = [(0.05, 0.08, 0.07), (0.92, 0.12, 0.05), (0.02, 0.55, 0.05), (0.97, 0.45, 0.08), (0.10, 0.93, 0.06),
             (0.88, 0.85, 0.05), (0.50, 0.02, 0.06), (0.62, 0.97, 0.04), (0.35, 0.40, 0.04), (0.70, 0.30, 0.05)]
    for sx, sy, d in stars:
        dot(slide, x + w * sx, y + h * sy, d, "white", alpha=70)
    dot(slide, x + w * 0.88, y + h * 0.05, 0.42, "sky", alpha=55)
    dot(slide, x + w * 0.0, y + h * 0.75, 0.3, "pink", alpha=60)
    dot(slide, x + w * 0.95, y + h * 0.65, 0.22, "purple", alpha=70)


def card(slide, x, y, w, h, shadow=0, fill="surface"):
    return box(slide, x, y, w, h, fill=fill, line="hairline", radius=0.17, shadow=shadow)


def notes(slide, body: str):
    slide.notes_slide.notes_text_frame.text = body


def patch_theme_fonts(prs):
    """테마 글꼴(차트·기본 텍스트)도 Paperlogy로 바꾼다."""
    theme = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(theme.blob)
    for tag, family in (("a:majorFont", F["bold"]), ("a:minorFont", F["regular"])):
        node = root.find(".//" + qn(tag))
        for child in ("a:latin", "a:ea", "a:cs"):
            node.find(qn(child)).set("typeface", family)
    theme._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


# ------------------------------------------------------------------ 슬라이드
def s_title(prs):
    s = new_slide(prs, "indigo")
    logo_placeholder(s, X0, 0.42)
    text(s, X0, 1.95, 8.2, 0.35, ("코디세이 AI 네이티브 Final-Project  ·  금융", "eyebrow", "sky"))
    text(s, X0, 2.45, 8.4, 2.1, [("끊어진 사다리,", "display1", "white"), ("첫 칸부터 다시 놓는다", "display1", "white")],
         line=1.05)
    text(s, X0, 4.62, 8.2, 0.4, ("사회초년생을 위한 개인별 자산관리 AI 코치, 사다리", "h3", "night_text"))
    b1 = box(s, X0, 5.4, 1.9, 0.52, fill="primary", radius=0.26)
    label_in(b1, ("라이브 데모", "button", "white"))
    b2 = box(s, X0 + 2.05, 5.4, 1.9, 0.52, fill="surface", radius=0.26, shadow=1)
    label_in(b2, ("GitHub 저장소", "button", "ink"))
    text(s, X0, 6.75, 8, 0.3, ("팀 ○○○  ·  팀장 ○○○  ·  2026. 11.", "body_sm", "night_muted"))
    constellation(s, 8.9, 0.75, 3.8, 6.0)
    notes(s, "안녕하세요, 팀 ○○○입니다. 저희는 사회초년생이 첫 월급부터 자산을 만들고, 그 계획을 끝까지 지키도록 돕는 AI 코치 '사다리'를 만들었습니다. "
             "'끊어진 계층 사다리의 첫 칸을 다시 놓는다'는 뜻을 담았습니다.")


def s_problem(prs, page):
    s = new_slide(prs)
    header(s, "01  문제 정의", "같은 청년 안에서도 자산 격차는 40배",
           "출발선의 차이가 '정보'와 '습관'의 차이로 굳어지고 있습니다")
    card(s, X0, 2.2, 5.6, 4.45, shadow=1)
    pill(s, X0 + 0.35, 2.55, "청년층 순자산 격차", style="eyebrow", h=0.3, pad=0.14)
    text(s, X0 + 0.35, 2.95, 5.0, 1.45, ("40.4배", "stat_xl"))
    text(s, X0 + 0.35, 4.5, 5.0, 0.4, ("순자산 상위 20%와 하위 20% 사이", "h3"))
    text(s, X0 + 0.35, 5.0, 4.9, 0.9, ("같은 세대 안에서도 출발선이 이만큼 벌어져 있고, "
                                      "정부의 자산형성 지원 가입률마저 계층에 따라 차이가 납니다.", "body"), line=1.3)
    text(s, X0 + 0.35, 6.2, 4.9, 0.3, ("출처: 국회입법조사처 보고서 (2026.7)", "caption"))
    rows = [
        ("purple", "62.6점", "20대 금융이해력. 전체 평균 65.7점보다 낮고, 재무목표·재무점검 점수는 2년 새 급락",
         "한국은행·금감원 (2024)"),
        ("pink", "39.3%", "청년정책을 이용하지 않은 이유 2위: \"관심은 있지만 찾는 방법을 몰라서\"",
         "청년정책 인식 조사 (2026.9, n=281)"),
        ("orange", "15.9%", "청년도약계좌 중도해지율. 2023년 말 8.2%에서 두 배 가까이 증가",
         "국회입법조사처 (2026.7)"),
    ]
    x = 6.5
    for i, (color, stat, desc, src) in enumerate(rows):
        y = 2.2 + i * 1.55
        card(s, x, y, CW - (x - X0), 1.35)
        dot(s, x + 0.3, y + 0.5, 0.34, color)
        text(s, x + 0.85, y + 0.28, 2.0, 0.8, ("" + stat, "stat"), anchor="middle")
        text(s, x + 2.85, y + 0.22, 3.2, 0.72, (desc, "body_sm", "ink"), line=1.2)
        text(s, x + 2.85, y + 0.98, 3.2, 0.25, (src, "caption", "faint"))
    footer(s, "문제 정의", page)
    notes(s, "먼저 문제입니다. 같은 청년층 안에서도 순자산 상위 20%와 하위 20%의 격차는 40배가 넘습니다. "
             "그런데 이 격차는 정책이 없어서 생기는 게 아닙니다. 20대의 금융이해력은 가장 낮고, 정책은 찾기 어렵고, 가입해도 중간에 해지합니다. "
             "판단하고, 실행하고, 유지하는 능력의 차이가 격차를 만든다고 봤습니다.")


def s_chart(prs, page):
    s = new_slide(prs)
    header(s, "01  문제 정의", "20대의 목표와 점검 능력이 무너졌다",
           "한국은행·금감원 「2024 전국민 금융이해력 조사」 20대 세부 점수 (100점 만점)")
    card(s, X0, 2.2, 7.4, 4.45)
    data = CategoryChartData()
    data.categories = ["재무목표 세우기", "재무 점검하기"]
    data.add_series("2022년", (48.0, 55.8))
    data.add_series("2024년", (36.1, 33.2))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(X0 + 0.3), Inches(2.4), Inches(6.8),
                            Inches(4.1), data)
    ch = gf.chart
    ch.font.size = Pt(14)
    ch.font.name = F["regular"]
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(13)
    plot = ch.plots[0]
    plot.gap_width = 70
    plot.overlap = -12
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = "0.0"
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(16)
    for series, color in zip(plot.series, ("C9C4BE", C["orange"]), strict=True):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(color)
    va = ch.value_axis
    va.minimum_scale, va.maximum_scale = 0, 70
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = _rgb("hairline")
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(12)
    va.tick_labels.font.color.rgb = _rgb("faint")
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(15)
    ca.format.line.color.rgb = _rgb("hairline")

    x = 8.3
    card(s, x, 2.2, CW - (x - X0), 4.45, shadow=1)
    pill(s, x + 0.35, 2.55, "사다리의 답", style="eyebrow", h=0.3, pad=0.14)
    text(s, x + 0.35, 3.0, 3.8, 0.8, ("목표는 함께 세우고, 점검은 매주 함께", "h3"), line=1.15)
    items = [("purple", "재무목표 36.1점 → 월급 배분 플랜", "5분 진단이면 무엇을 얼마씩 모을지 계산 엔진이 정해줍니다"),
             ("teal", "재무점검 33.2점 → 주간 체크인", "목표 진척과 적금 납입 위험을 매주 알려주고 약속을 기억합니다")]
    for i, (color, head, desc) in enumerate(items):
        y = 4.05 + i * 1.2
        dot(s, x + 0.35, y + 0.08, 0.22, color)
        text(s, x + 0.72, y, 3.35, 0.35, (head, "title"))
        text(s, x + 0.72, y + 0.38, 3.35, 0.7, (desc, "body_sm"), line=1.2)
    footer(s, "문제 정의", page)
    notes(s, "특히 눈에 띄는 건 20대의 '재무목표'와 '재무점검' 점수입니다. 2년 만에 각각 12점, 22점이 떨어졌습니다. "
             "그래서 사다리는 목표는 계산 엔진이 함께 세우고, 점검은 매주 코치가 함께하는 구조로 설계했습니다.")


def s_pain(prs, page):
    s = new_slide(prs)
    header(s, "01  문제 정의", "청년이 실제로 겪는 네 가지 문제", "각 문제마다 사다리의 기능이 하나씩 대응합니다")
    items = [
        ("purple", "P1", "무엇부터 해야 할지 모른다", "월급은 들어오는데 배분 기준이 없다. 20대 재무목표 36.1점",
         "5분 진단 → 우선순위가 정해진 월급 배분 플랜"),
        ("pink", "P2", "나에게 맞는 혜택을 못 찾는다", "정책은 많지만 내가 대상인지, 얼마를 받는지 모른다",
         "자격 판정 + 근거 검색으로 받을 수 있는 혜택과 예상액"),
        ("orange", "P3", "가입해도 유지하지 못한다", "3년, 5년짜리 적금을 중간에 깬다. 도약계좌 해지율 15.9%",
         "약속을 기억하는 AI 코치와 주간 체크인"),
        ("teal", "P4", "모을 돈 자체가 없다", "정책상품은 '본인 저축'을 전제로 한다. 저축 여력이 먼저다",
         "지출 영향 계산과 소비 분석으로 저축 여력 만들기"),
    ]
    w, h = 5.965, 2.1
    for i, (color, tag, head, evidence, answer) in enumerate(items):
        x = X0 + (i % 2) * (w + 0.2)
        y = 2.2 + (i // 2) * (h + 0.25)
        card(s, x, y, w, h)
        num_badge(s, x + 0.32, y + 0.32, 0.56, tag, color, "ink" if color in ("purple", "pink") else "white")
        text(s, x + 1.1, y + 0.36, w - 1.4, 0.45, (head, "h3"))
        text(s, x + 1.1, y + 0.84, w - 1.4, 0.45, (evidence, "body_sm"), line=1.2)
        dot(s, x + 1.12, y + 1.43, 0.13, color)
        text(s, x + 1.38, y + 1.32, w - 1.68, 0.66, (answer, "title"), line=1.15)
    footer(s, "문제 정의", page)
    notes(s, "청년들의 문제를 네 가지로 정리했습니다. 무엇부터 할지 모르고, 맞는 혜택을 못 찾고, 가입해도 유지를 못 하고, 애초에 모을 돈이 없다는 것입니다. "
             "사다리의 기능은 이 네 가지에 하나씩 대응하도록 설계했습니다.")


def s_compare(prs, page):
    s = new_slide(prs)
    header(s, "02  해결", "조회가 아니라, 실행과 유지까지", "기존 서비스와 사다리의 역할 비교")
    cols = [
        ("자산관리 앱", "토스·뱅크샐러드 등", ["계좌 통합 조회와 소비 분석", "거래 데이터",
                                      "청년 정책과 재무 계획의 연결은 부차적", "상품 제휴 수익 구조"], False),
        ("정책 포털", "온통청년·보조금24 등", ["정책 검색과 자격 조회", "나이·지역 등 인적 조건",
                                        "내 현금흐름은 반영되지 않고 우선순위 판단은 사용자 몫", "중립"], False),
        ("사다리", "청년 자산관리 AI 코치", ["현금흐름 × 정책 × 목표를 잇는 실행 계획과 유지 코칭",
                                       "재무 상황 + 인적 조건 + 목표·약속(기억)",
                                       "계좌 자동 연동 대신 직접 입력해 민감정보 최소화", "상품 판매·중개를 하지 않음"], True),
    ]
    labels = ["핵심 가치", "개인화 기준", "한계", "중립성"]
    w = 3.9
    for i, (name, ex, values, featured) in enumerate(cols):
        x = X0 + i * (w + 0.215)
        card(s, x, 2.2, w, 4.45, shadow=2 if featured else 0)
        text(s, x + 0.3, 2.4, w - 0.6, 0.42, (name, "h2"))
        text(s, x + 0.3, 2.84, w - 0.6, 0.3, (ex, "caption"))
        if featured:
            pill(s, x + w - 1.45, 2.45, "우리의 접근", style="eyebrow", h=0.3, pad=0.14)
        for j, (lab, val) in enumerate(zip(labels, values, strict=True)):
            y = 3.32 + j * 0.83
            box(s, x + 0.3, y - 0.08, w - 0.6, 0.01, fill="hairline")
            text(s, x + 0.3, y, w - 0.6, 0.25, (lab, "caption", "faint"))
            text(s, x + 0.3, y + 0.25, w - 0.6, 0.5, (val, "body_sm", "ink" if featured else "ink2"), line=1.15)
    footer(s, "해결", page)
    notes(s, "기존 자산관리 앱은 조회와 분석에, 정책 포털은 목록 제공에 머뭅니다. "
             "사다리는 내 현금흐름과 청년 정책, 그리고 목표를 하나의 실행 계획으로 묶고, 그 계획이 무너지지 않도록 곁에서 코칭합니다. "
             "계좌를 자동으로 연결하지 않는 대신 민감정보를 받지 않고, 상품을 팔지 않아 중립적입니다.")


def s_flow(prs, page):
    s = new_slide(prs)
    header(s, "02  해결", "진단 → 설계 → 매칭 → 실행 → 유지",
           "가상 인물 '지우' (27세, 중소기업 1년 차, 월 실수령 245만 원)의 여정")
    steps = [
        ("purple", "진단", "질문 카드로 소득·지출·대출·목표를 파악", "5분, 질문 10~12개"),
        ("pink", "설계", "계산 엔진이 월급을 필수지출·생활비·비상금·저축으로 배분", "저축 65만·비상금 19만"),
        ("orange", "매칭", "받을 수 있는 청년 혜택과 예상 수혜액을 출처와 함께", "기여금 약 199만 원"),
        ("teal", "실행", "\"이거 사도 돼?\" 같은 질문에 내 숫자로 답하고 목표를 조정", "30만 쓰면 저축 16만"),
        ("green", "유지", "약속을 기억하고 매주 목표 진척과 납입 위험을 점검", "배달 주 2회 약속 기억"),
    ]
    w, h, gap = 2.2, 2.85, 0.2825
    for i, (color, name, desc, example) in enumerate(steps):
        x = X0 + i * (w + gap)
        card(s, x, 2.2, w, h)
        num_badge(s, x + 0.25, 2.45, 0.5, str(i + 1), color, "ink" if color in ("purple", "pink") else "white")
        text(s, x + 0.25, 3.1, w - 0.5, 0.4, (name, "h2"))
        text(s, x + 0.25, 3.6, w - 0.5, 0.85, (desc, "body_sm"), line=1.2)
        box(s, x + 0.25, 4.43, w - 0.5, 0.01, fill="hairline")
        text(s, x + 0.2, 4.52, w - 0.4, 0.35, (example, "chip"))
        if i < len(steps) - 1:
            arrow(s, x + w + 0.03, 3.575, x + w + gap - 0.03, 3.575, color="faint", width=1.25)
    card(s, X0, 5.27, CW, 1.38, shadow=1)
    dot(s, X0 + 0.35, 5.65, 0.34, "sky")
    text(s, X0 + 0.9, 5.47, 3.2, 0.4, ("지우의 경우", "h3"))
    text(s, X0 + 0.9, 5.9, CW - 1.3, 0.7,
         ("목표 '2년 안에 전세보증금 3천만 원'은 지금 속도(월 46만 원)로는 66개월이 걸립니다. "
          "코치는 기간 조정, 청년미래적금 기여금 활용, 새는 지출 줄이기를 숫자와 함께 제안합니다.", "body"), line=1.25)
    footer(s, "해결", page)
    notes(s, "사다리는 진단, 설계, 매칭, 실행, 유지의 다섯 단계로 동작합니다. 가상 인물 지우로 예를 들면, 5분 진단 후 월 65만 원 저축 플랜을 받고, "
             "청년미래적금으로 3년간 약 199만 원의 기여금을 받을 수 있다는 걸 알게 됩니다. "
             "그리고 '이번 달 30만 원 더 써도 될까?'라고 물으면, 목표 저축이 16만 원으로 줄어든다는 걸 숫자로 보여줍니다.")


def s_screens(prs, page):
    s = new_slide(prs)
    header(s, "02  해결", "실제로 동작하는 MVP", "모바일 웹·PWA  ·  예시 계정으로 로그인 없이 체험  ·  화면은 모의 AI 모드 기준")
    shots = [("00_onboarding.jpg", "진단: 질문 카드"), ("01_home.jpg", "홈: 이번 달 요약"), ("02_plan.jpg", "월급 배분"),
             ("03_benefit.jpg", "혜택 자격 판정"), ("04_coach.jpg", "AI 코치")]
    h = 4.15
    w = h * 750 / 1624
    step = (CW - 0.18 - w) / (len(shots) - 1)
    for i, (file, cap) in enumerate(shots):
        x = X0 + 0.09 + i * step
        phone(s, SCREENS / file, x, 2.18, h)
        text(s, x - 0.3, 6.5, w + 0.6, 0.3, (cap, "title"), align="center")
    footer(s, "해결", page)
    notes(s, "실제로 동작하는 화면입니다. 진단은 한 화면에 질문 하나씩, 금액은 만 원 단위로 받습니다. "
             "홈에서 이번 달 저축과 목표를 보고, 플랜 화면에서 배분 이유를 확인합니다. "
             "혜택 화면은 자격 조건을 하나씩 판정해 보여주고, 코치는 계산 엔진의 숫자로 답합니다. (라이브 데모로 전환)")


def s_divider(prs, page, number, title, subtitle, section):
    s = new_slide(prs, "indigo")
    pill(s, X0, 0.42, number, style="eyebrow", h=0.32, pad=0.16)
    text(s, X0, 0.86, 8.3, 0.9, (title, "display2", "white"))
    text(s, X0, 1.85, 7.8, 0.9, (subtitle, "h3", "night_text"), line=1.3)
    constellation(s, 8.9, 0.75, 3.8, 6.0)
    for i, (color, word) in enumerate([("purple", "계산은 코드"), ("teal", "근거는 공식 문서"), ("sky", "대화와 설명은 AI")]):
        y = 3.2 + i * 0.78
        dot(s, X0, y + 0.1, 0.28, color)
        text(s, X0 + 0.5, y, 6, 0.45, (word, "h2", "white"))
    text(s, X0, 5.95, 7.8, 0.6, ("이어서: AI를 쓰는 곳과 쓰지 않는 곳  ·  코치 Agent  ·  RAG  ·  안전장치와 품질 측정",
                                  "body_sm", "night_muted"), line=1.3)
    footer(s, section, page, dark=True)
    notes(s, "이제 AI를 어떻게 설계했는지 말씀드리겠습니다. 원칙은 하나입니다. 틀리면 안 되는 것은 코드로, 근거는 공식 문서에서, 대화와 설명만 AI가 합니다.")


def s_ai_where(prs, page):
    s = new_slide(prs)
    header(s, "03  AI 설계", "AI를 쓰는 곳과 쓰지 않는 곳", "틀리면 안 되는 판단은 코드로, 해석과 대화는 AI로")
    cols = [
        ("AI가 맡는 일", "green", ["진단의 목표 문장 해석 (\"2년 안에 전세 3천\")", "질문 의도 파악과 도구 선택",
                                "플랜 설명: 계산 결과만 보고 쉽게 풀어 쓰기", "다음 대화에 쓸 약속·선호 추출",
                                "(예정) 영수증 사진 인식"]),
        ("코드가 맡는 일", "orange", ["금액·만기액·정부기여금 계산", "정책 자격 판정: 나이·소득·주거·모집기간",
                                   "정책 문서 검색 (BM25)", "민감정보 가리기·투자권유 차단",
                                   "답변 속 금액이 계산 결과와 같은지 검증"]),
    ]
    w = 5.965
    for i, (head, color, items) in enumerate(cols):
        x = X0 + i * (w + 0.2)
        card(s, x, 2.2, w, 3.3)
        dot(s, x + 0.35, 2.5, 0.3, color)
        text(s, x + 0.8, 2.44, w - 1.1, 0.42, (head, "h3"))
        for j, item in enumerate(items):
            y = 3.02 + j * 0.47
            dot(s, x + 0.43, y + 0.1, 0.12, "faint")
            text(s, x + 0.72, y, w - 1.0, 0.4, (item, "body"))
    card(s, X0, 5.72, CW, 0.93, shadow=1)
    text(s, X0 + 0.35, 5.84, CW - 0.7, 0.7,
         [[("왜 이렇게 나눴나?  ", "title"),
          ("LLM은 산술을 틀릴 수 있고, 정책은 자주 바뀝니다(청년도약계좌 신규가입 종료 2025.12, "
           "청년미래적금 출시 2026.6). 그래서 숫자는 계산 엔진이, 정책은 최신 공식 문서가 근거가 됩니다.", "body_sm")]],
         anchor="middle", line=1.25)
    footer(s, "AI 설계", page)
    notes(s, "AI를 쓰는 곳과 쓰지 않는 곳을 명확히 나눴습니다. 금액 계산과 자격 판정은 틀리면 안 되기 때문에 결정론적인 코드가 맡습니다. "
             "AI는 목표 문장을 해석하고, 질문에 맞는 도구를 고르고, 결과를 쉽게 설명하는 역할만 합니다.")


def s_agent(prs, page):
    s = new_slide(prs)
    header(s, "03  AI 설계", "코치 Agent: 질문을 이해하고 도구를 고른다",
           "Claude 도구 호출 루프  ·  최대 6턴  ·  SSE 스트리밍  ·  API 키가 없으면 같은 도구를 쓰는 모의 코치")
    cy = 3.7
    nodes = [(X0, 1.25, "질문", None), (2.2, 2.05, "입력 가드레일", "민감정보 마스킹\n투자권유 요청 탐지"),
             (4.62, 2.5, "코치 Agent", "Claude Opus 5.5\n도구 선택 · 설명"), (7.49, 2.05, "출력 가드레일", "권유 문구 교체\n근거 없는 금액 표시"),
             (9.91, 1.3, "답변 + 출처", None), (11.58, 1.15, "장기 메모리", None)]
    for i, (x, w, title_, desc) in enumerate(nodes):
        big = title_ == "코치 Agent"
        h = 1.55 if big else (1.2 if desc else 0.85)
        y = cy - h / 2
        card(s, x, y, w, h, shadow=2 if big else 0)
        if desc:
            text(s, x + 0.18, y + 0.18, w - 0.36, 0.38, (title_, "h3" if big else "title"), align="center")
            text(s, x + 0.12, y + 0.6 + (0.12 if big else 0), w - 0.24, 0.6, (desc, "caption"), align="center", line=1.15)
        else:
            text(s, x + 0.05, y, w - 0.1, h, (title_, "title"), align="center", anchor="middle")
        if i < len(nodes) - 1:
            nx = nodes[i + 1][0]
            arrow(s, x + w + 0.03, cy, nx - 0.03, cy, color="muted", width=1.25)
    arrow(s, 5.87, cy + 0.8, 5.87, 5.08, color="muted", width=1.25, both=True)
    card(s, X0, 5.12, CW, 1.53)
    text(s, X0 + 0.3, 5.22, CW - 0.6, 0.3, ("도구 7개. 모든 금액은 도구(계산 엔진)가 계산하고, AI는 결과만 설명합니다", "caption"))
    tools = [("재무 조회", "get_my_finances"), ("저축 계산", "simulate_savings"), ("지출 영향", "check_extra_spending"),
             ("정책 검색", "search_youth_policies"), ("정책 상세", "get_policy_detail"), ("기억 저장", "save_memory"),
             ("목표 설정", "set_goal")]
    tw, gap = 1.6, 0.143
    for i, (ko, code) in enumerate(tools):
        x = X0 + 0.3 + i * (tw + gap)
        b = box(s, x, 5.6, tw, 0.82, fill="canvas", line="hairline", radius=0.1)
        label_in(b, [(ko, "title"), (code, "code")], margin=0.03)
    footer(s, "AI 설계", page)
    notes(s, "코치의 구조입니다. 질문이 들어오면 먼저 입력 가드레일이 민감정보를 가리고 투자 권유 요청을 걸러냅니다. "
             "코치 Agent는 질문에 맞는 도구를 골라 호출하는데, 도구는 7개이고 모든 금액은 도구가 계산합니다. "
             "답변은 출력 가드레일에서 권유 문구와 근거 없는 금액을 한 번 더 확인한 뒤, 출처와 함께 스트리밍으로 전달되고, 약속은 장기 메모리에 저장됩니다.")


def s_rag(prs, page):
    s = new_slide(prs)
    header(s, "03  AI 설계", "근거 있는 답: 검색 + 규칙 판정 + 출처",
           "정책 문서 검색(BM25)과 규칙 기반 자격 판정을 결합한 RAG")
    steps = [("purple", "정책 데이터", "7개 정책 · 출처 URL · 기준일\n(20개 이상으로 확장 중)"),
             ("pink", "BM25 검색", "한글 음절 단위 검색\n정책명이 나오면 우선"),
             ("orange", "자격 판정", "나이·소득·주거·재직·모집 기간\n규칙 코드가 판정"),
             ("teal", "근거 있는 답변", "근거 문장 + 출처 + 기준일\n모르면 공식 기관 확인 안내")]
    w, gap = 2.8, 0.31
    for i, (color, head, desc) in enumerate(steps):
        x = X0 + i * (w + gap)
        card(s, x, 2.2, w, 1.8)
        dot(s, x + 0.28, 2.47, 0.26, color)
        text(s, x + 0.66, 2.41, w - 0.9, 0.4, (head, "h3"))
        text(s, x + 0.28, 2.95, w - 0.4, 0.9, (desc, "body_sm"), line=1.25)
        if i < len(steps) - 1:
            arrow(s, x + w + 0.04, 3.1, x + w + gap - 0.04, 3.1, color="faint", width=1.25)
    card(s, X0, 4.25, 7.9, 2.4, shadow=1)
    text(s, X0 + 0.35, 4.48, 3, 0.4, ("청년미래적금", "h3"))
    pill(s, X0 + 2.2, 4.5, "조건 확인 필요", style="chip", fill="EAF3FD", color="ink", line=None, h=0.3, pad=0.14)
    text(s, X0 + 0.35, 4.97, 7.2, 0.35, ("월 46만 원 납입 시 3년간 정부기여금 약 198만 7,200원 (우대형 가정)", "title"))
    checks = [("green", "나이", "만 19~34세 대상 (현재 만 27세)"), ("green", "소득", "총급여 7,500만 원 이하 (입력 연봉 3,400만 원)"),
              ("sky", "추가 확인", "가구 기준 중위소득 200% 이하 (우대형은 150% 이하)")]
    for i, (color, lab, desc) in enumerate(checks):
        y = 5.45 + i * 0.33
        dot(s, X0 + 0.38, y + 0.07, 0.14, color)
        text(s, X0 + 0.65, y, 1.2, 0.3, (lab, "chip"))
        text(s, X0 + 1.8, y, 5.8, 0.3, (desc, "body_sm"))
    text(s, X0 + 0.35, 6.4, 7.2, 0.22, ("출처: KB Think, 토스뱅크 안내 · 기준일 2026-09-30 · 공식 공고 재검수 중", "caption", "faint"))
    x = 8.75
    card(s, x, 4.25, CW - (x - X0), 2.4)
    text(s, x + 0.3, 4.45, 3.4, 0.4, ("판정 결과 4단계", "h3"))
    legend = [("green", "받을 수 있어요", "모든 조건 충족"), ("sky", "조건 확인 필요", "가구소득 등 확인"),
              ("orange", "지금은 신청 불가", "모집·가입 종료"), ("faint", "해당 없음", "조건 불충족")]
    for i, (color, lab, desc) in enumerate(legend):
        y = 4.98 + i * 0.4
        dot(s, x + 0.32, y + 0.07, 0.16, color)
        text(s, x + 0.6, y, 1.5, 0.3, (lab, "chip"))
        text(s, x + 2.1, y, 1.65, 0.3, (desc, "caption"))
    footer(s, "AI 설계", page)
    notes(s, "RAG는 네 단계입니다. 출처와 기준일이 붙은 정책 데이터를 BM25로 검색하고, 자격은 LLM이 아니라 규칙 코드가 판정합니다. "
             "가구소득처럼 사다리가 알 수 없는 조건은 '조건 확인 필요'로 솔직하게 표시하고, 모집 기간이 끝난 정책은 '지금은 신청 불가'로 보여줍니다.")


def s_quality(prs, page):
    s = new_slide(prs)
    header(s, "03  AI 설계", "안전장치와 품질 측정",
           "가드레일 4겹과 자동 평가로 AI 품질을 숫자로 관리합니다 (평가셋 40·30문항으로 확장 중)")
    w = 6.0
    card(s, X0, 2.2, w, 4.45)
    text(s, X0 + 0.35, 2.45, w - 0.7, 0.4, ("가드레일 4겹", "h3"))
    guards = [("입력", "민감정보 마스킹", "주민·계좌·카드·전화번호를 AI 전송 전에 가림 (조사가 붙어도)"),
              ("입력", "투자권유 요청 탐지", "\"삼성전자 사도 돼?\" → 투자 원칙 안내로 대체"),
              ("출력", "권유 문구 차단", "\"매수하세요\" 등이 섞이면 안전 안내문으로 교체"),
              ("출력", "수치 검증", "도구 결과에 없는 금액은 검토 대상으로 표시")]
    for i, (stage, head, desc) in enumerate(guards):
        y = 3.0 + i * 0.8
        num_badge(s, X0 + 0.35, y + 0.05, 0.5, str(i + 1), "purple" if stage == "입력" else "teal",
                  "ink" if stage == "입력" else "white")
        text(s, X0 + 1.05, y, w - 1.35, 0.35, [[(f"{stage}  ", "chip", "faint"), (head, "title")]])
        text(s, X0 + 1.05, y + 0.36, w - 1.35, 0.35, (desc, "body_sm"))
    text(s, X0 + 0.35, 6.25, w - 0.7, 0.28, ("모든 답변에 'AI 답변' 표시  ·  사용자당 하루 50회 호출 상한", "caption"))
    stats = [("79개", "자동 테스트 통과", "계산 엔진·자격 판정·가드레일·전체 흐름"),
             ("100%", "정책 검색 정답률", "hit@1, 골든셋 12문항"),
             ("100%", "투자권유 요청 탐지", "레드팀 10문항"),
             ("0원", "테스트·CI 비용", "API 키 없이 도는 모의 AI 모드")]
    x0, tw, th = 6.85, 2.84, 2.12
    for i, (num, head, desc) in enumerate(stats):
        x = x0 + (i % 2) * (tw + 0.2)
        y = 2.2 + (i // 2) * (th + 0.21)
        card(s, x, y, tw, th, shadow=1)
        text(s, x + 0.3, y + 0.25, tw - 0.5, 0.75, (num, "stat"))
        text(s, x + 0.3, y + 1.08, tw - 0.5, 0.35, (head, "title"))
        text(s, x + 0.3, y + 1.45, tw - 0.5, 0.55, (desc, "caption"), line=1.15)
    footer(s, "AI 설계", page)
    notes(s, "안전장치는 입력 두 겹, 출력 두 겹으로 총 네 겹입니다. 그리고 AI 품질을 느낌이 아니라 숫자로 관리합니다. "
             "자동 테스트 79개, 정책 검색 정답률, 투자 권유 탐지율을 CI에서 매번 측정하고, 평가셋은 각각 40문항과 30문항으로 늘리고 있습니다.")


def s_stack(prs, page):
    s = new_slide(prs)
    header(s, "04  구현", "단순하게, 설명할 수 있게", "모든 기술 결정은 ADR(의사결정 기록) 8건으로 남겼습니다")
    w = 6.3
    card(s, X0, 2.2, w, 4.45)
    user = box(s, X0 + 1.55, 2.45, 3.2, 0.6, fill="canvas", line="hairline", radius=0.3)
    label_in(user, ("사용자 · 모바일 웹 / PWA", "title"))
    arrow(s, X0 + 3.15, 3.07, X0 + 3.15, 3.37, color="muted", width=1.25)
    box(s, X0 + 0.35, 3.4, w - 0.7, 1.65, fill="surface", line="muted", radius=0.14, line_w=1.0)
    text(s, X0 + 0.55, 3.5, w - 1.1, 0.3, ("FastAPI 단일 서비스 (Render)", "title"))
    inner = ["화면\nJinja2", "API · SSE\n스트리밍", "코치 Agent\n도구 호출", "계산 엔진\n정책 검색"]
    iw = (w - 0.7 - 0.4 - 3 * 0.12) / 4
    for i, lab in enumerate(inner):
        b = box(s, X0 + 0.55 + i * (iw + 0.12), 3.92, iw, 0.95, fill="canvas", line=None, radius=0.1)
        label_in(b, (lab, "chip"), margin=0.03)
    outs = [("Claude API", "purple"), ("SQLite → PostgreSQL", "teal"), ("n8n 자동화", "orange")]
    ow = (w - 0.7 - 2 * 0.2) / 3
    for i, (lab, color) in enumerate(outs):
        x = X0 + 0.35 + i * (ow + 0.2)
        arrow(s, x + ow / 2, 5.07, x + ow / 2, 5.5, color="muted", width=1.25)
        b = box(s, x, 5.53, ow, 0.75, fill="surface", line="hairline", radius=0.1)
        label_in(b, (lab, "chip"), margin=0.03)
        dot(s, x + 0.12, 5.6, 0.12, color)
    x = X0 + w + 0.25
    rw = CW - (x - X0)
    card(s, x, 2.2, rw, 4.45)
    text(s, x + 0.3, 2.42, rw - 0.6, 0.4, ("기술 결정 기록 (ADR)", "h3"))
    adrs = [("단일 서비스", "Python 하나로 개발, 서비스 1개로 배포"),
            ("모의 AI 모드", "API 키 없이 전체 흐름 동작, 테스트 비용 0원"),
            ("BM25 검색", "수백 청크 규모엔 벡터 DB가 필요 없음"),
            ("무료 배포", "DB는 늦게 만들어 30일 만료를 발표 뒤로"),
            ("계산은 코드", "LLM의 산술 오류를 구조적으로 차단"),
            ("세션 인증", "외부 인증 서비스 없이 scrypt + 서명 쿠키"),
            ("SDK 루프", "프레임워크 없이 60줄, 팀 누구나 설명 가능"),
            ("디자인 시스템", "강조색 하나, 한 화면에 질문 하나")]
    for i, (lab, why) in enumerate(adrs):
        y = 2.95 + i * 0.46
        box(s, x + 0.3, y - 0.07, rw - 0.6, 0.01, fill="hairline")
        text(s, x + 0.3, y + 0.02, 1.55, 0.35, (lab, "chip"))
        text(s, x + 1.85, y + 0.02, rw - 2.15, 0.4, (why, "body_sm"))
    footer(s, "구현", page)
    notes(s, "구현은 '단순하게, 설명할 수 있게'를 원칙으로 했습니다. 화면과 API를 FastAPI 하나로 만들고, Agent도 프레임워크 없이 짧은 루프로 작성해서 팀원 누구나 설명할 수 있습니다. "
             "기획 초안에서 바뀐 결정은 모두 ADR로 남겨, 왜 바꿨는지 추적할 수 있습니다.")


def s_team(prs, page):
    s = new_slide(prs)
    header(s, "05  협업", "5명이 나눠 맡는 26개의 이슈", "기능 단위 오너십  ·  PR 교차 리뷰  ·  CI 자동 검증")
    roles = [("purple", "팀장·PM", ["협업 체계·이슈 등록", "피드백 → GitHub Issue", "사용자 테스트·발표"]),
             ("pink", "프론트엔드", ["목표 대시보드", "소비 분석 화면", "영수증 촬영·PWA 설치"]),
             ("orange", "백엔드·인프라", ["거래내역 분석 API", "PostgreSQL 전환", "보안 강화·저축 기록"]),
             ("teal", "AI Agent", ["실제 Claude 연결·모델 비교", "영수증 인식 (멀티모달)", "레드팀·응답 품질 평가"]),
             ("green", "RAG·자동화", ["정책 데이터 검수·확장", "온통청년·금감원 API", "n8n 주간 체크인"])]
    w, gap = 2.3, 0.1575
    for i, (color, name, issues) in enumerate(roles):
        x = X0 + i * (w + gap)
        card(s, x, 2.2, w, 3.2)
        band = box(s, x, 2.2, w, 0.72, fill=color, radius=0.17, shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
        label_in(band, (name, "h3", "ink" if color in ("purple", "pink") else "white"), align="left", margin=0.25)
        text(s, x + 0.25, 3.08, w - 0.5, 0.3, ("담당 ○○○", "caption"))
        for j, issue in enumerate(issues):
            y = 3.5 + j * 0.58
            dot(s, x + 0.27, y + 0.1, 0.1, color)
            text(s, x + 0.48, y, w - 0.68, 0.55, (issue, "body_sm", "ink"), line=1.15)
    card(s, X0, 5.62, CW, 1.03, shadow=1)
    flow = ["이슈 (FR-ID)", "브랜치", "PR · 리뷰 1인 이상", "CI: 테스트 · AI 평가", "머지 · 자동 배포"]
    widths = [1.75, 1.2, 2.1, 2.3, 1.95]
    total = sum(widths) + 0.45 * (len(flow) - 1)
    x = X0 + (CW - total) / 2
    for i, (lab, fw) in enumerate(zip(flow, widths, strict=True)):
        pill(s, x, 5.93, lab, style="chip", fill="canvas", w=fw, h=0.42)
        if i < len(flow) - 1:
            arrow(s, x + fw + 0.07, 6.14, x + fw + 0.38, 6.14, color="faint", width=1.25)
        x += fw + 0.45
    footer(s, "협업", page)
    notes(s, "협업 구조입니다. 필수 기능으로 MVP 기반을 먼저 만들고, 나머지 기능을 역할별 이슈 26개로 나눠 다섯 명이 균형 있게 맡았습니다. "
             "모든 작업은 이슈에서 시작해 PR 리뷰와 CI 검증을 거쳐야 머지되고, 머지되면 자동으로 배포됩니다. 이 기록이 모두 GitHub에 남습니다.")


def s_result(prs, page):
    s = new_slide(prs)
    header(s, "06  검증", "사용자 테스트 결과",
           "만 19~34세 청년 10명 · 1주 실사용 (10/24~11/1)  ·  [테스트 후 입력]")
    metrics = [("SUS 사용성 점수", "목표 70점 이상"), ("새로 알게 된 혜택", "목표 1인당 2개 이상"),
               ("재사용 의향", "목표 60% 이상"), ("과제 성공률", "목표 80% 이상")]
    w = 2.875
    for i, (head, target) in enumerate(metrics):
        x = X0 + i * (w + 0.21)
        card(s, x, 2.2, w, 1.9)
        text(s, x + 0.3, 2.42, w - 0.6, 0.3, (head, "title"))
        text(s, x + 0.3, 2.78, w - 0.6, 0.75, ("—", "stat", "faint"))
        text(s, x + 0.3, 3.62, w - 0.6, 0.3, (target, "caption"))
    card(s, X0, 4.35, CW, 2.3, shadow=1)
    text(s, X0 + 0.35, 4.55, 5, 0.4, ("주요 피드백과 개선 (Before → After)", "h3"))
    cols = [("피드백 (익명 인용)", 0.35, 4.3), ("조치", 4.8, 3.2), ("Before → After", 8.15, 3.6)]
    for lab, cx, cw in cols:
        text(s, X0 + cx, 5.05, cw, 0.3, (lab, "chip", "muted"))
    for r in range(3):
        y = 5.45 + r * 0.38
        box(s, X0 + 0.35, y - 0.05, CW - 0.7, 0.01, fill="hairline")
        for _, cx, cw in cols:
            text(s, X0 + cx, y, cw, 0.3, ("[테스트 후 입력]", "body_sm", "faint"))
    footer(s, "검증", page)
    notes(s, "[테스트 후 작성] 사용자 테스트는 10명을 대상으로 1주일간 진행했습니다. 핵심 지표와, 피드백을 받아 실제로 무엇을 바꿨는지를 Before/After로 보여드립니다.")


def s_impact(prs, page):
    s = new_slide(prs)
    header(s, "07  기대효과", "개인의 첫 칸에서 정책의 실효성까지", "기대효과와 발전 방향")
    cols = [("purple", "개인", ["첫 월급부터 '무엇을 먼저' 할지 분명한 실행 계획",
                               "놓치던 청년 혜택과 예상 수혜액 발견", "주간 코칭으로 적금 중도해지 예방"]),
            ("teal", "사회", ["정보 격차가 자산 격차로 이어지는 고리 완화",
                             "청년 자산형성 정책의 '가입 → 유지' 실효성 제고", "20대 재무목표·점검 역량 향상"]),
            ("orange", "비즈니스", ["B2G·B2B: 청년센터·대학·기업에 재무코칭 도구 공급",
                                   "B2C: 기본 무료 + 심화 리포트 구독", "상품 판매·중개는 하지 않음 (중립성·규제)"])]
    w = 3.9
    for i, (color, head, items) in enumerate(cols):
        x = X0 + i * (w + 0.215)
        card(s, x, 2.2, w, 3.3)
        dot(s, x + 0.35, 2.52, 0.3, color)
        text(s, x + 0.8, 2.46, w - 1.1, 0.42, (head, "h2"))
        for j, item in enumerate(items):
            y = 3.12 + j * 0.74
            dot(s, x + 0.4, y + 0.1, 0.11, "faint")
            text(s, x + 0.65, y, w - 0.95, 0.7, (item, "body_sm", "ink"), line=1.2)
    card(s, X0, 5.72, CW, 0.93, shadow=1)
    road = [("지금", "MVP · 사용자 테스트"), ("다음", "청년센터·대학 파일럿"), ("확장", "마이데이터 제휴·대상 확대")]
    x = X0 + 0.35
    for i, (when, what) in enumerate(road):
        pill(s, x, 6.0, when, style="eyebrow", h=0.36, pad=0.16)
        text(s, x + 0.8, 5.98, 3.2, 0.4, (what, "title"), anchor="middle")
        if i < len(road) - 1:
            arrow(s, x + 3.35, 6.18, x + 3.75, 6.18, color="faint", width=1.25)
        x += 3.95
    footer(s, "기대효과", page)
    notes(s, "개인에게는 첫 월급부터의 실행 계획을, 사회적으로는 정보 격차가 자산 격차로 이어지는 고리를 끊는 데 기여합니다. "
             "비즈니스는 공공성이 큰 청년센터와 대학을 첫 시장으로 보고 있으며, 투자자문이나 상품 중개는 하지 않아 규제 리스크를 줄였습니다.")


def s_closing(prs, page):
    s = new_slide(prs, "indigo")
    logo_placeholder(s, X0, 0.42)
    text(s, X0, 2.1, 8.2, 0.35, ("감사합니다", "eyebrow", "sky"))
    text(s, X0, 2.55, 8.2, 1.9, [("첫 칸을", "display1", "white"), ("함께 놓아요", "display1", "white")], line=1.05)
    text(s, X0, 4.62, 8.2, 0.4, ("예시 계정으로 로그인 없이 바로 써 볼 수 있습니다", "h3", "night_text"))
    b1 = box(s, X0, 5.4, 2.1, 0.52, fill="primary", radius=0.26)
    label_in(b1, ("사다리 체험하기", "button", "white"))
    b2 = box(s, X0 + 2.25, 5.4, 1.9, 0.52, fill="surface", radius=0.26, shadow=1)
    label_in(b2, ("GitHub 저장소", "button", "ink"))
    text(s, X0, 6.4, 8, 0.3, ("Q&A  ·  배포 URL과 저장소 주소는 제출 전에 기재", "body_sm", "night_muted"))
    constellation(s, 8.9, 0.75, 3.8, 6.0)
    notes(s, "사다리는 청년의 첫 칸을 함께 놓는 코치입니다. 예시 계정으로 지금 바로 써 보실 수 있습니다. 감사합니다. 질문 받겠습니다.")


def build() -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    patch_theme_fonts(prs)
    s_title(prs)
    page = 2
    for fn in (s_problem, s_chart, s_pain, s_compare, s_flow, s_screens):
        fn(prs, page)
        page += 1
    s_divider(prs, page, "03", "AI는 꼭 필요한 곳에만", "틀리면 안 되는 것은 코드로, 근거는 공식 문서에서, 대화와 설명은 AI가 맡습니다", "AI 설계")
    page += 1
    for fn in (s_ai_where, s_agent, s_rag, s_quality, s_stack, s_team, s_result, s_impact):
        fn(prs, page)
        page += 1
    s_closing(prs, page)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    print("초안 생성:", build())
