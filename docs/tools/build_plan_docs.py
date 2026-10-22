"""기획안 Markdown(docs/01_기획안.md)을 한글(HWPX)과 워드(DOCX) 파일로 만든다.

- HWPX: 코디세이 공식 양식(docs/tools/templates/plan_form.hwpx, [붙임4] 상세기획서 양식)의 칸에 내용을 채운다.
  양식의 글꼴·표·여백은 그대로 두고, 2쪽 안에 들어가도록 작은 글자 스타일만 추가한다.
- DOCX: 같은 표 구조를 python-docx로 그린다.

사용법:
    python docs/tools/build_plan_docs.py              # docs/01_기획안.hwpx, docs/01_기획안.docx
필요 패키지: python-docx (DOCX만). HWPX는 표준 라이브러리만 쓴다.
PDF는 build_plan_pdf.py로 따로 만든다. 세 형식 모두 01_기획안.md 하나에서 나온다.
"""

from __future__ import annotations

import re
import sys
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from xml.sax.saxutils import escape

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_plan_pdf import LIST_RE, parse  # noqa: E402  같은 Markdown 해석 규칙을 쓴다

DOCS = Path(__file__).resolve().parent.parent
SRC = DOCS / "01_기획안.md"
TEMPLATE = Path(__file__).resolve().parent / "templates" / "plan_form.hwpx"

INLINE_RE = re.compile(r"\*\*(.+?)\*\*|`(.+?)`")
BULLETS = ["• ", "- ", "· "]

# 표 머리글로 열 너비 비율을 정한다 (없으면 글자 수로 추정)
COLUMN_RATIOS = {
    ("역할", "담당자", "담당 업무"): [0.22, 0.10, 0.68],
    ("핵심 기능", "내용", "AI 기술"): [0.19, 0.61, 0.20],
}


# ------------------------------------------------------------------ Markdown → 블록
def runs(text: str) -> list[tuple[str, bool]]:
    """'**굵게**'와 '`코드`'를 (글, 굵게 여부) 조각으로 나눈다."""
    out, pos = [], 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], False))
        out.append((m.group(1), True) if m.group(1) is not None else (m.group(2), False))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], False))
    return out


def blocks(lines: list[str]) -> list[tuple]:
    """("p", 글) / ("li", 깊이, 기호, 글) / ("h4", 글) / ("table", 머리글, 행들)"""
    out: list[tuple] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
        elif line.lstrip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            out.append(("table", rows[0], rows[2:]))
        elif m := LIST_RE.match(line):
            out.append(("li", len(m.group(1)) // 2, m.group(2), m.group(3)))
            i += 1
        elif line.startswith("### "):
            out.append(("h4", line[4:].strip()))
            i += 1
        else:
            para = []
            while (i < len(lines) and lines[i].strip() and not lines[i].lstrip().startswith("|")
                   and not LIST_RE.match(lines[i])):
                para.append(lines[i].strip())
                i += 1
            out.append(("p", " ".join(para)))
    return out


def marker(depth: int, sign: str) -> str:
    return f"{sign} " if sign[0].isdigit() else BULLETS[min(depth, len(BULLETS) - 1)]


def column_ratios(head: list[str], rows: list[list[str]]) -> list[float]:
    plain = tuple(re.sub(r"\*\*|`", "", h) for h in head)
    if plain in COLUMN_RATIOS:
        return COLUMN_RATIOS[plain]
    weights = [max(8, min(60, max(len(r[c]) for r in [head, *rows]))) for c in range(len(head))]
    return [w / sum(weights) for w in weights]


# ------------------------------------------------------------------ HWPX
# 양식에 추가하는 스타일 id (양식에는 charPr 0~22, paraPr 0~23이 있다)
CH_BODY, CH_BOLD, CH_TBL, CH_TBL_BOLD = 23, 24, 25, 26
PP_BODY, PP_LIST1, PP_LIST2, PP_NUM, PP_HEAD, PP_TBL, PP_TBL_CENTER, PP_VALUE = range(24, 32)


def _char_pr(base: str, new_id: int, height: int) -> str:
    return re.sub(r'id="\d+" height="\d+"', f'id="{new_id}" height="{height}"', base, count=1)


def _para_pr(base: str, new_id: int, align: str, left: int = 0, intent: int = 0, prev: int = 0, nxt: int = 0,
             line: int = 140) -> str:
    """양식의 paraPr를 복제한다. 한글은 hp:case 값의 두 배를 hp:default에 기록한다."""
    def margin(k: int) -> str:
        vals = {"intent": intent, "left": left, "right": 0, "prev": prev, "next": nxt}
        inner = "".join(f'<hc:{n} value="{v * k}" unit="HWPUNIT"/>' for n, v in vals.items())
        return f'<hh:margin>{inner}</hh:margin><hh:lineSpacing type="PERCENT" value="{line}" unit="HWPUNIT"/>'

    switch = ('<hp:switch><hp:case hp:required-namespace="http://www.hancom.co.kr/hwpml/2016/HwpUnitChar">'
              f"{margin(1)}</hp:case><hp:default>{margin(2)}</hp:default></hp:switch>")
    out = re.sub(r'<hh:paraPr id="\d+"', f'<hh:paraPr id="{new_id}"', base, count=1)
    out = re.sub(r'<hh:align horizontal="\w+"', f'<hh:align horizontal="{align}"', out, count=1)
    return re.sub(r"<hp:switch>.*?</hp:switch>", switch, out, count=1, flags=re.S)


def patch_header(header: str) -> str:
    char_normal = re.search(r'<hh:charPr id="18".*?</hh:charPr>', header, re.S).group(0)
    char_bold = re.search(r'<hh:charPr id="19".*?</hh:charPr>', header, re.S).group(0)
    new_chars = [_char_pr(char_normal, CH_BODY, 900), _char_pr(char_bold, CH_BOLD, 900),
                 _char_pr(char_normal, CH_TBL, 850), _char_pr(char_bold, CH_TBL_BOLD, 850)]
    base = re.search(r'<hh:paraPr id="18".*?</hh:paraPr>', header, re.S).group(0)
    new_paras = [
        _para_pr(base, PP_BODY, "JUSTIFY", nxt=120),
        _para_pr(base, PP_LIST1, "JUSTIFY", left=200, intent=-900, nxt=60),
        _para_pr(base, PP_LIST2, "JUSTIFY", left=1100, intent=-700, nxt=40),
        _para_pr(base, PP_NUM, "JUSTIFY", left=200, intent=-1100, nxt=60),
        _para_pr(base, PP_HEAD, "LEFT", prev=200, nxt=60),
        _para_pr(base, PP_TBL, "LEFT", line=125),
        _para_pr(base, PP_TBL_CENTER, "CENTER", line=125),
        _para_pr(base, PP_VALUE, "CENTER", line=130),
    ]

    def append(xml: str, tag: str, items: list[str]) -> str:
        count = int(re.search(rf'<hh:{tag} itemCnt="(\d+)"', xml).group(1))
        xml = xml.replace(f'<hh:{tag} itemCnt="{count}"', f'<hh:{tag} itemCnt="{count + len(items)}"', 1)
        return xml.replace(f"</hh:{tag}>", "".join(items) + f"</hh:{tag}>", 1)

    header = append(header, "charProperties", new_chars)
    return append(header, "paraProperties", new_paras)


def hp_p(parts: list[tuple[str, bool]], para: int, normal: int, bold: int, prefix: str = "") -> str:
    body = f'<hp:run charPrIDRef="{normal}"><hp:t>{escape(prefix)}</hp:t></hp:run>' if prefix else ""
    for text, is_bold in parts:
        body += f'<hp:run charPrIDRef="{bold if is_bold else normal}"><hp:t>{escape(text)}</hp:t></hp:run>'
    if not body:
        body = f'<hp:run charPrIDRef="{normal}"/>'
    return f'<hp:p id="0" paraPrIDRef="{para}" styleIDRef="0" pageBreak="0" columnBreak="0" merged="0">{body}</hp:p>'


class HwpxTables:
    def __init__(self) -> None:
        self.next_id = 1991800000
        self.z = 20

    def xml(self, head: list[str], rows: list[list[str]], width: int) -> str:
        ratios = column_ratios(head, rows)
        widths = [round(width * r) for r in ratios]
        widths[-1] = width - sum(widths[:-1])
        row_h = 1300
        all_rows = [head, *rows]
        trs = []
        for r, cells in enumerate(all_rows):
            tcs = []
            for c, text in enumerate(cells):
                if r == 0:
                    para = hp_p([(re.sub(r"\*\*", "", text), True)], PP_TBL_CENTER, CH_TBL, CH_TBL_BOLD)
                else:
                    para = hp_p(runs(text), PP_TBL, CH_TBL, CH_TBL_BOLD)
                tcs.append(
                    f'<hp:tc name="" header="0" hasMargin="0" protect="0" editable="0" dirty="0" '
                    f'borderFillIDRef="{10 if r == 0 else 4}"><hp:subList id="" textDirection="HORIZONTAL" '
                    f'lineWrap="BREAK" vertAlign="CENTER" linkListIDRef="0" linkListNextIDRef="0" textWidth="0" '
                    f'textHeight="0" hasTextRef="0" hasNumRef="0">{para}</hp:subList>'
                    f'<hp:cellAddr colAddr="{c}" rowAddr="{r}"/><hp:cellSpan colSpan="1" rowSpan="1"/>'
                    f'<hp:cellSz width="{widths[c]}" height="{row_h}"/>'
                    f'<hp:cellMargin left="283" right="283" top="85" bottom="85"/></hp:tc>')
            trs.append("<hp:tr>" + "".join(tcs) + "</hp:tr>")
        self.next_id += 1
        self.z += 1
        tbl = (f'<hp:tbl id="{self.next_id}" zOrder="{self.z}" numberingType="TABLE" textWrap="TOP_AND_BOTTOM" '
               f'textFlow="BOTH_SIDES" lock="0" dropcapstyle="None" pageBreak="CELL" repeatHeader="1" '
               f'rowCnt="{len(all_rows)}" colCnt="{len(head)}" cellSpacing="0" borderFillIDRef="4" noAdjust="0">'
               f'<hp:sz width="{width}" widthRelTo="ABSOLUTE" height="{row_h * len(all_rows)}" '
               f'heightRelTo="ABSOLUTE" protect="0"/><hp:pos treatAsChar="1" affectLSpacing="0" flowWithText="1" '
               f'allowOverlap="0" holdAnchorAndSO="0" vertRelTo="PARA" horzRelTo="PARA" vertAlign="TOP" '
               f'horzAlign="LEFT" vertOffset="0" horzOffset="0"/><hp:outMargin left="0" right="0" top="0" '
               f'bottom="0"/><hp:inMargin left="283" right="283" top="85" bottom="85"/>'
               + "".join(trs) + "</hp:tbl>")
        return (f'<hp:p id="0" paraPrIDRef="{PP_BODY}" styleIDRef="0" pageBreak="0" columnBreak="0" merged="0">'
                f'<hp:run charPrIDRef="{CH_BODY}">{tbl}<hp:t/></hp:run></hp:p>')


def section_paras(lines: list[str], tables: HwpxTables, width: int) -> str:
    out = []
    for block in blocks(lines):
        kind = block[0]
        if kind == "p":
            out.append(hp_p(runs(block[1]), PP_BODY, CH_BODY, CH_BOLD))
        elif kind == "h4":
            out.append(hp_p([(block[1], True)], PP_HEAD, CH_BODY, CH_BOLD))
        elif kind == "li":
            _, depth, sign, text = block
            para = PP_NUM if sign[0].isdigit() else (PP_LIST1 if depth == 0 else PP_LIST2)
            out.append(hp_p(runs(text), para, CH_BODY, CH_BOLD, prefix=marker(depth, sign)))
        elif kind == "table":
            out.append(tables.xml(block[1], block[2], width))
    return "".join(out)


def build_hwpx(meta: dict[str, str], sections: list[tuple[str, list[str]]], dst: Path) -> None:
    src = zipfile.ZipFile(TEMPLATE)
    files = {info.filename: src.read(info.filename) for info in src.infolist()}
    infos = src.infolist()

    header = patch_header(files["Contents/header.xml"].decode("utf-8"))
    section = files["Contents/section0.xml"].decode("utf-8")

    values = {(1, 1): meta.get("도메인", ""), (1, 3): meta.get("팀명", ""),
              (2, 1): meta.get("프로젝트명", ""), (2, 3): meta.get("팀장명", "")}
    content_rows = {"기획의도": 3, "역할분담": 4, "프로젝트 내용": 5, "기대효과 및 활용방안": 6}
    by_row = {content_rows[title]: lines for title, lines in sections if title in content_rows}
    tables = HwpxTables()

    def fill_cell(m: re.Match) -> str:
        tc = m.group(0)
        row = int(re.search(r'rowAddr="(\d+)"', tc).group(1))
        col = int(re.search(r'colAddr="(\d+)"', tc).group(1))
        if (row, col) in values:
            char = re.search(r'<hp:run charPrIDRef="(\d+)"', tc).group(1)
            inner = hp_p([(values[(row, col)], False)], PP_VALUE, int(char), int(char))
        elif col == 1 and row in by_row:
            width = int(re.search(r'<hp:cellSz width="(\d+)"', tc).group(1)) - 1020  # 셀 좌우 여백
            inner = section_paras(by_row[row], tables, width)
            tc = tc.replace('vertAlign="CENTER"', 'vertAlign="TOP"', 1)
        else:
            return tc
        return re.sub(r"(<hp:subList [^>]*>).*?(</hp:subList>)", lambda s: s.group(1) + inner + s.group(2),
                      tc, count=1, flags=re.S)

    # 바깥 양식 표(제목 아래)의 칸만 채운다. 양식 안에는 중첩 표가 없으므로 비탐욕 일치로 칸을 나눌 수 있다
    start = section.index('<hp:tbl id="1991748639"')
    end = section.index("</hp:tbl>", start) + len("</hp:tbl>")
    table = re.sub(r"<hp:tc .*?</hp:tc>", fill_cell, section[start:end], flags=re.S)
    section = section[:start] + table + section[end:]
    section = section.replace("<hp:t>※ 2쪽 이내로 자유롭게 작성</hp:t>", "<hp:t/>")

    title = f"{meta.get('프로젝트명', '')} Final-Project 기획안"
    hpf = files["Contents/content.hpf"].decode("utf-8")
    hpf = re.sub(r"<opf:title>.*?</opf:title>", f"<opf:title>{escape(title)}</opf:title>", hpf)
    hpf = re.sub(r'(<opf:meta name="ModifiedDate" content="text">)[^<]*',
                 r"\g<1>" + datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ"), hpf)
    preview = "\r\n".join(["Final-Project 프로젝트 기획안", title,
                           *[f"<{t}>" for t, _ in sections]])

    files["Contents/header.xml"] = header.encode("utf-8")
    files["Contents/section0.xml"] = section.encode("utf-8")
    files["Contents/content.hpf"] = hpf.encode("utf-8")
    files["Preview/PrvText.txt"] = preview.encode("utf-8")

    # mimetype이 맨 앞, 압축하지 않음 (양식과 같은 순서·압축 방식 유지)
    with zipfile.ZipFile(dst, "w") as out:
        for info in infos:
            out.writestr(info, files[info.filename], compress_type=info.compress_type)


class CellWriter:
    """표 칸 안에 문단과 표를 차례로 붙인다. 칸에 처음 있는 빈 문단과 표 뒤에 Word가 요구하는 빈 문단을 다시 쓴다."""

    def __init__(self, cell) -> None:
        self.cell = cell
        self.spare = cell.paragraphs[0]

    def par(self):
        if self.spare is not None:
            par, self.spare = self.spare, None
            return par
        return self.cell.add_paragraph()

    def table(self, rows: int, cols: int):
        from docx.shared import Pt

        starts_cell = len(self.cell._tc.xpath("./w:p")) == 1 and not self.cell._tc.xpath("./w:tbl")
        if self.spare is not None and starts_cell:
            # 칸이 표로 시작하면 앞의 빈 문단을 지운다 (표 뒤에는 python-docx가 빈 문단을 붙인다)
            self.spare._element.getparent().remove(self.spare._element)
        table = self.cell.add_table(rows=rows, cols=cols)
        self.spare = self.cell.paragraphs[-1]
        self.spare.paragraph_format.space_after = Pt(2)
        return table


# ------------------------------------------------------------------ DOCX
def build_docx(meta: dict[str, str], sections: list[tuple[str, list[str]]], dst: Path) -> None:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    font = "맑은 고딕"
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(15)
    sec.top_margin, sec.bottom_margin = Mm(13), Mm(12)
    content_w = 210 - 30

    normal = doc.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(9)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = Pt(14)  # 맑은 고딕의 '한 줄'은 글자 크기의 1.7배라 고정값으로 맞춘다
    ppr = normal.element.get_or_add_pPr()
    for tag in ("w:autoSpaceDE", "w:autoSpaceDN"):  # '40 배'처럼 숫자와 한글 사이가 벌어지지 않게
        el = OxmlElement(tag)
        el.set(qn("w:val"), "0")
        ppr.append(el)

    def shade(cell, color: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), color)
        tc_pr.append(shd)

    def borders(cell, **sides) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        tcb = OxmlElement("w:tcBorders")
        for side, (val, size) in sides.items():
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:val"), val)
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), "000000")
            tcb.append(el)
        tc_pr.append(tcb)

    def cell_margins(table, top=50, bottom=50, left=110, right=110) -> None:
        tbl_pr = table._tbl.tblPr
        mar = OxmlElement("w:tblCellMar")
        for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:w"), str(v))
            el.set(qn("w:type"), "dxa")
            mar.append(el)
        tbl_pr.append(mar)

    def write(par, parts, size=9.0, bold_all=False, color=None) -> None:
        if size > 9.5:  # 제목·라벨처럼 큰 글자는 고정 줄 간격을 글자 크기에 맞춘다
            par.paragraph_format.line_spacing = Pt(round(size * 1.6))
        for text, is_bold in parts:
            r = par.add_run(text)
            r.bold = is_bold or bold_all
            r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)

    def set_widths(table, widths_mm) -> None:
        table.autofit = False
        for row in table.rows:
            for cell, w in zip(row.cells, widths_mm, strict=False):
                cell.width = Mm(w)

    # 머리띠: [붙임1] Final-Project 프로젝트 기획안
    band = doc.add_table(rows=1, cols=2)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_widths(band, [20, content_w - 20])
    b0, b1 = band.rows[0].cells
    shade(b0, "342DBE")
    borders(b0, top=("single", 12), bottom=("single", 12), left=("single", 12), right=("single", 12))
    borders(b1, top=("single", 12), bottom=("single", 12), left=("nil", 0), right=("nil", 0))
    for cell in (b0, b1):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    b0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    write(b0.paragraphs[0], [("붙임1", True)], size=12, color="FFFFFF")
    write(b1.paragraphs[0], [("Final-Project 프로젝트 기획안", True)], size=13)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # 본 양식 표
    form = doc.add_table(rows=7, cols=4)
    form.style = "Table Grid"
    form.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_margins(form, top=45, bottom=45, left=120, right=120)
    widths = [w / 47624 * content_w for w in (9359, 18415, 7661, 12189)]
    set_widths(form, widths)

    title = form.cell(0, 0).merge(form.cell(0, 3))
    shade(title, "D9D9D9")
    title.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    write(title.paragraphs[0], [("코디세이 AI 네이티브 응용 학습 Final-Project 기획안", True)], size=13)

    pairs = [(1, 0, "도메인", meta.get("도메인", "")), (1, 2, "팀 명", meta.get("팀명", "")),
             (2, 0, "프로젝트명", meta.get("프로젝트명", "")), (2, 2, "팀장 명", meta.get("팀장명", ""))]
    for r, c, key, val in pairs:
        k, v = form.cell(r, c), form.cell(r, c + 1)
        shade(k, "F2F2F2")
        for cell, text, bold in ((k, key, True), (v, val, False)):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            write(cell.paragraphs[0], [(text, bold)], size=10.5)

    inner_w = widths[1] + widths[2] + widths[3] - 5
    order = ["기획의도", "역할분담", "프로젝트 내용", "기대효과 및 활용방안"]
    content = dict(sections)
    for i, name in enumerate(order):
        r = 3 + i
        label = form.cell(r, 0)
        label.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        label.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        text = name.replace(" 및 ", " 및\n") if " 및 " in name else name
        for j, line in enumerate(text.split("\n")):
            par = label.paragraphs[0] if j == 0 else label.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            write(par, [(line, True)], size=10)
        body = CellWriter(form.cell(r, 1).merge(form.cell(r, 3)))
        body.cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        for block in blocks(content.get(name, [])):
            kind = block[0]
            if kind in ("p", "h4"):
                par = body.par()
                par.paragraph_format.space_after = Pt(3)
                if kind == "h4":
                    par.paragraph_format.space_before = Pt(4)
                write(par, runs(block[1]), bold_all=kind == "h4")
            elif kind == "li":
                _, depth, sign, text = block
                par = body.par()
                fmt = par.paragraph_format
                hang = 4.0 if sign[0].isdigit() else 3.2
                fmt.left_indent = Mm(1 + depth * 4.5 + hang)
                fmt.first_line_indent = Mm(-hang)
                fmt.space_after = Pt(1.5)
                write(par, [(marker(depth, sign), False), *runs(text)])
            elif kind == "table":
                head, rows = block[1], block[2]
                t = body.table(len(rows) + 1, len(head))
                t.style = "Table Grid"
                cell_margins(t, top=30, bottom=30, left=80, right=80)
                set_widths(t, [inner_w * x for x in column_ratios(head, rows)])
                for row in t.rows:  # 행이 쪽 경계에서 쪼개지지 않게
                    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))  # 다음 쪽에 머리글 반복
                for c, text in enumerate(head):
                    cell = t.cell(0, c)
                    shade(cell, "F2F2F2")
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cell.paragraphs[0].paragraph_format.keep_with_next = True  # 머리글만 쪽 끝에 남지 않게
                    write(cell.paragraphs[0], [(re.sub(r"\*\*", "", text), True)], size=8.5)
                for ri, row in enumerate(rows, start=1):
                    for c, text in enumerate(row):
                        cell = t.cell(ri, c)
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        cell.paragraphs[0].paragraph_format.line_spacing = Pt(12)
                        write(cell.paragraphs[0], runs(text), size=8.5)
    doc.core_properties.title = f"{meta.get('프로젝트명', '')} Final-Project 기획안"
    doc.core_properties.author = "사다리 팀"
    doc.save(dst)


def main() -> None:
    meta, sections = parse(SRC.read_text(encoding="utf-8"))
    hwpx = SRC.with_suffix(".hwpx")
    build_hwpx(meta, sections, hwpx)
    print(f"생성 완료: {hwpx}")
    try:
        import docx  # noqa: F401
    except ImportError:
        print("python-docx가 없어 DOCX는 건너뜁니다 (pip install python-docx)")
        return
    out = SRC.with_suffix(".docx")
    build_docx(meta, sections, out)
    print(f"생성 완료: {out}")


if __name__ == "__main__":
    main()
